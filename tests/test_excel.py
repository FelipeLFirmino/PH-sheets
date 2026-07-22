"""
Testes de integração para salvar_excel_estilizado().

Objetivo: garantir que a geração do arquivo Excel não levanta KeyError ao
referenciar chaves do dicionário COL — bug que ocorreu quando P_ATC foi
renomeado para P_ATC_PED/P_ATC_PDV mas a lista number_format ainda usava
o nome antigo.

Estes testes não validam os valores das células (isso é responsabilidade de
test_calcular.py), mas verificam que a função executa sem erros e produz um
arquivo .xlsx válido.
"""
import os
import tempfile
import pytest
from tests.conftest import make_row
from core.processador import salvar_excel_estilizado, COL, F_RATEIO
from openpyxl import load_workbook


@pytest.fixture
def P():
    return {
        'mult':         2.0,
        'frete':        0.10,
        'desp':         0.10,
        'cred':         0.04,
        'fed':          0.0913,
        'icm':          0.21,
        'cartao':       0.04,
        'mult_atc':     1.3,
        'desc_atc':     0.15,
        'desc_atc_pdv': 0.10,
    }


def _make_full_row(**kwargs):
    """
    make_row com os campos extras exigidos pelo Excel gerado em salvar_excel_estilizado().
    Inclui todos os campos que gerar_tabela() popula na lista de rows mas que make_row()
    não possui (pois make_row é voltado apenas para _calcular()).
    """
    defaults = {
        'desc':      'PRODUTO TESTE',
        'ref':       'REF001',
        'sku':       'SKU001',
        'nf':        '001',
        'cred_pct':  0.04,
        'qtd_emb':   1,
        'p_atual':   0.0,
        'p_sys_raw': 0.0,   # preço bruto do sistema (auditoria)
    }
    base = make_row()
    base.update(defaults)
    base.update(kwargs)
    # Derivar valores raw (por unidade comercial, antes da divisão por qtd_emb)
    # a partir dos valores já definidos — espelha o que gerar_tabela() faz.
    qtd_emb = base.get('qtd_emb', 1) or 1
    for key in ('nf_u', 'st_u', 'ant_u', 'ipi_u'):
        raw_key = key + '_raw'
        if raw_key not in base:
            base[raw_key] = base.get(key, 0.0) * qtd_emb
    return base


class TestExcelGeracaoSemErros:
    """Garante que salvar_excel_estilizado não levanta exceções para inputs válidos."""

    def test_gera_sem_key_error_produto_normal(self, P):
        """Produto sem ST/ANT não deve levantar KeyError em nenhuma coluna."""
        row = _make_full_row(nf_u=10.0)
        dados = ([row], P, '0001')
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            salvar_excel_estilizado(dados, path)
            assert os.path.exists(path)
            wb = load_workbook(path)
            assert 'Precificação' in wb.sheetnames
        finally:
            os.unlink(path)

    def test_gera_sem_key_error_produto_com_st(self, P):
        """Produto com ST não deve levantar KeyError."""
        row = _make_full_row(nf_u=10.0, st_u=3.0, tem_st=True)
        dados = ([row], P, '0001')
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            salvar_excel_estilizado(dados, path)
            assert os.path.exists(path)
        finally:
            os.unlink(path)

    def test_gera_sem_key_error_produto_com_ant(self, P):
        """Produto com ANT não deve levantar KeyError."""
        row = _make_full_row(nf_u=10.0, ant_u=2.0, tem_ant=True)
        dados = ([row], P, '0001')
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            salvar_excel_estilizado(dados, path)
            assert os.path.exists(path)
        finally:
            os.unlink(path)

    def test_gera_sem_key_error_multiplos_produtos(self, P):
        """Múltiplos produtos com tipos variados não devem levantar KeyError."""
        rows = [
            _make_full_row(nf_u=10.0),
            _make_full_row(nf_u=20.0, st_u=5.0, tem_st=True),
            _make_full_row(nf_u=15.0, ant_u=2.0, tem_ant=True),
            _make_full_row(nf_u=8.0, ipi_u=1.0),
            _make_full_row(nf_u=12.0, p_atual=25.0),
        ]
        dados = (rows, P, '0001')
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            salvar_excel_estilizado(dados, path)
            assert os.path.exists(path)
        finally:
            os.unlink(path)

    def test_produto_rateado_marca_desc_e_colore_st_ant(self, P):
        """
        Produto com 'rateado'=True (contagem XML != contagem API para a mesma
        descrição — ver parear_impostos_api) deve: (1) ganhar o prefixo de
        aviso na descrição e (2) ter ST_U/ANT_U pintados de vermelho (F_RATEIO)
        para conferência manual, independente de ser linha ST/ANT/normal.
        """
        row = _make_full_row(nf_u=10.0, ant_u=2.0, tem_ant=True, rateado=True,
                              desc='ABAJUR DECORATIVO CX12')
        dados = ([row], P, '0001')
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            salvar_excel_estilizado(dados, path)
            wb = load_workbook(path)
            ws = wb['Precificação']
            assert ws.cell(3, COL['DESC']).value.startswith('⚠ CONFERIR IMPOSTO')
            assert ws.cell(3, COL['ST_U']).fill.start_color.rgb  == '00FF0000'
            assert ws.cell(3, COL['ANT_U']).fill.start_color.rgb == '00FF0000'
        finally:
            os.unlink(path)

    def test_produto_nao_rateado_nao_pinta_vermelho(self, P):
        """Produto normal (rateado=False) não deve ter ST_U/ANT_U em vermelho."""
        row = _make_full_row(nf_u=10.0, ant_u=2.0, tem_ant=True, rateado=False)
        dados = ([row], P, '0001')
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            path = f.name
        try:
            salvar_excel_estilizado(dados, path)
            wb = load_workbook(path)
            ws = wb['Precificação']
            assert not ws.cell(3, COL['DESC']).value.startswith('⚠')
            assert ws.cell(3, COL['ST_U']).fill.start_color.rgb != '00FF0000'
        finally:
            os.unlink(path)

    def test_todas_chaves_col_existem(self):
        """
        Testa que todas as chaves referenciadas em COL são válidas.
        Previne regressões do tipo KeyError: 'P_ATC' — quando uma chave
        é renomeada no dicionário mas ainda referenciada em outro lugar.

        Se este teste falhar, procure usages de COL['NOME_ANTIGO'] no código.
        """
        chaves_esperadas = {
            'NF', 'DESC', 'REF', 'SKU', 'QTD',
            'NF_U', 'ST_U', 'ANT_U', 'IPI_U',
            'C_REAL', 'FRETE', 'DESP', 'CRED', 'CST', 'C_ENT',
            'FED', 'CARTAO', 'ICMS_S', 'C_SAIDA',
            'META', 'P_MIN', 'P_ATUAL', 'P_VAR', 'MARGEM',
            'NF_ATC', 'P_ATC_PED', 'P_ATC_PDV',
            'FED_ATC', 'CART_ATC', 'ICM_ATC',
            'C_SAIDA_ATC', 'MARGEM_ATC_PED', 'MARGEM_ATC_PDV',
            'P_PCT_ATC', 'P_COMPRA_PCT',
            'AUDIT_SYS', 'AUDIT_EMB', 'AUDIT_CRED',
        }
        assert chaves_esperadas == set(COL.keys()), (
            f"Chaves ausentes: {chaves_esperadas - set(COL.keys())}\n"
            f"Chaves inesperadas: {set(COL.keys()) - chaves_esperadas}"
        )

    def test_col_indices_unicos_e_sequenciais(self):
        """Nenhum índice de coluna pode se repetir — dois nomes na mesma coluna causam bugs silenciosos."""
        valores = list(COL.values())
        assert len(valores) == len(set(valores)), (
            f"Índices duplicados: { {k: v for k, v in COL.items() if valores.count(v) > 1} }"
        )
        assert sorted(valores) == list(range(1, len(valores) + 1)), (
            "Índices não são sequenciais começando em 1"
        )
